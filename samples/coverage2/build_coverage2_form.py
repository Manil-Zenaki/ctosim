"""Builds the second coverage XLSForm and its fictitious preload file.

The first coverage form (samples/coverage/) settled most of the SurveyCTO export layout.
This form targets what its exports could not show (spec 08, "Not observed in the real
exports") and the cases its three submissions left empty:

- a geopoint inside a repeat, and a select_multiple inside a nested repeat;
- a counted repeat nested in another repeat (name of its count column);
- the instanceName column (settings: instance_name);
- what comments, audio audit and speed violations fields hold (resp_name and hh_size have a
  minimum time, so that moving on too fast records speed violations);
- choice values with letters and hyphens in select_multiple dummy column names;
- a relevant select_multiple left blank (dummy cells);
- a top-level enumerator-controlled repeat with instances, and a refused consent.

All names and values are fictitious.

Run:  .venv/Scripts/python samples/coverage2/build_coverage2_form.py
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

SURVEY_HEADER = ["type", "name", "label", "hint", "appearance", "constraint", "constraint message",
                 "relevance", "required", "calculation", "repeat_count", "minimum_seconds"]

SURVEY = [
    # SurveyCTO metadata, as the online designer adds them, plus Collect-only audits
    {"type": "start", "name": "starttime"},
    {"type": "end", "name": "endtime"},
    {"type": "deviceid", "name": "deviceid"},
    {"type": "username", "name": "username"},
    {"type": "calculate", "name": "duration", "calculation": "duration()"},
    {"type": "text audit", "name": "text_audit", "appearance": "p=100"},
    {"type": "audio audit", "name": "audio_audit", "appearance": "p=100; s=0; d=10"},
    {"type": "speed violations count", "name": "sv_count", "appearance": "100"},
    {"type": "speed violations list", "name": "sv_list", "appearance": "100"},
    {"type": "comments", "name": "field_comments"},

    {"type": "select_one hhid", "name": "hhid", "label": "Household ID", "required": "yes"},
    {"type": "calculate", "name": "pl_head", "calculation": "pulldata('hh_preload', 'head_name', 'hhid', ${hhid})"},
    {"type": "select_one yesno", "name": "consent", "label": "Does the respondent agree to the interview?",
     "required": "yes"},

    {"type": "begin group", "name": "main", "label": "Household", "relevance": "${consent} = 1"},
    {"type": "text", "name": "resp_name", "label": "Respondent's first name (fictitious; answer in under 10 seconds)",
     "required": "yes", "minimum_seconds": "10"},
    {"type": "select_multiple assets", "name": "assets", "label": "Assets owned",
     "constraint": "not(selected(., '-99') and count-selected(.) > 1)",
     "constraint message": "'Don't know' cannot be combined"},
    {"type": "select_multiple transport", "name": "transport",
     "label": "Means of transport used last month (choice values with letters and hyphens)"},
    {"type": "integer", "name": "hh_size",
     "label": "Number of household members (1 to 3 for this test; answer in under 10 seconds)",
     "constraint": ". >= 1 and . <= 3", "required": "yes", "minimum_seconds": "10"},

    # Counted repeat holding a geopoint, a nested counted repeat and a nested
    # enumerator-controlled repeat with a select_multiple
    {"type": "begin repeat", "name": "member", "label": "Member", "repeat_count": "${hh_size}"},
    {"type": "text", "name": "member_name", "label": "Member's first name (fictitious)", "required": "yes"},
    {"type": "geopoint", "name": "member_gps", "label": "Where is this member now? (any location)"},
    {"type": "select_multiple activities", "name": "member_activities",
     "label": "Activities (may be left blank)"},
    {"type": "integer", "name": "n_children", "label": "Number of children of this member (0 to 2)",
     "constraint": ". >= 0 and . <= 2", "required": "yes"},
    {"type": "begin repeat", "name": "child", "label": "Child", "repeat_count": "${n_children}"},
    {"type": "text", "name": "child_name", "label": "Child's first name (fictitious)"},
    {"type": "end repeat", "name": "child"},
    {"type": "begin repeat", "name": "job", "label": "Job (add 0 to 2)"},
    {"type": "text", "name": "job_title", "label": "Job title"},
    {"type": "select_multiple days", "name": "job_days", "label": "Days worked"},
    {"type": "end repeat", "name": "job"},
    {"type": "end repeat", "name": "member"},

    # Top-level enumerator-controlled repeat
    {"type": "begin repeat", "name": "plot", "label": "Plot (add 1 or 2)"},
    {"type": "select_one crop", "name": "plot_crop", "label": "Main crop"},
    {"type": "geopoint", "name": "plot_gps", "label": "Plot location (any location)"},
    {"type": "end repeat", "name": "plot"},
    {"type": "end group", "name": "main"},
]

CHOICES = [
    ("yesno", "1", "Yes"), ("yesno", "0", "No"),
    ("assets", "1", "Radio"), ("assets", "2", "Bicycle"), ("assets", "96", "Other"),
    ("assets", "-99", "Don't know"),
    ("transport", "bus", "Bus"), ("transport", "moto-taxi", "Motorbike taxi"),
    ("transport", "x2", "Other (code with a digit)"),
    ("activities", "1", "Farming"), ("activities", "2", "Trade"), ("activities", "-97", "None"),
    ("days", "1", "Monday"), ("days", "6", "Saturday"), ("days", "-99", "Don't know"),
    ("crop", "1", "Maize"), ("crop", "2", "Teff"),
    ("hhid", "H001", "H001"), ("hhid", "H002", "H002"), ("hhid", "H003", "H003"),
]

PRELOAD = [
    ("hhid", "head_name"),
    ("H001", "Abebe (fictitious)"),
    ("H002", "Chaltu (fictitious)"),
    ("H003", "Dawit (fictitious)"),
]


def main() -> None:
    wb = Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(SURVEY_HEADER)
    for row in SURVEY:
        survey.append([row.get(column) for column in SURVEY_HEADER])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "value", "label"])
    for row in CHOICES:
        choices.append(list(row))

    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "default_language", "instance_name"])
    settings.append(["ctosim coverage test 2", "ctosim_coverage2", "2610100005", "english",
                     "concat(${hhid}, ' - ', ${resp_name})"])
    wb.save(HERE / "ctosim_coverage2.xlsx")

    with open(HERE / "hh_preload.csv", "w", encoding="utf-8", newline="\n") as f:
        for row in PRELOAD:
            f.write(",".join(row) + "\n")


if __name__ == "__main__":
    main()
