"""Builds the coverage XLSForm and its fictitious preload file.

The form exercises every field type and structure that the SurveyCTO export must
represent (metadata, groups, repeats, nested repeats, select_multiple with negative
codes, geopoint, dates and times, media, comments, text audit, preloads), so that a
handful of fictitious test submissions exported from SurveyCTO show the exact export
layout. No real data is involved.

Run:  .venv/Scripts/python samples/coverage/build_coverage_form.py
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

SURVEY_HEADER = ["type", "name", "label", "hint", "default", "appearance", "constraint",
                 "constraint message", "relevance", "required", "read only", "calculation",
                 "repeat_count", "choice_filter", "disabled"]

# Each row lists values for SURVEY_HEADER by column name.
SURVEY = [
    # SurveyCTO metadata, as the online designer adds them
    {"type": "start", "name": "starttime"},
    {"type": "end", "name": "endtime"},
    {"type": "deviceid", "name": "deviceid"},
    {"type": "subscriberid", "name": "subscriberid"},
    {"type": "simserial", "name": "simid"},
    {"type": "phonenumber", "name": "devicephonenum"},
    {"type": "username", "name": "username"},
    {"type": "calculate", "name": "duration", "calculation": "duration()"},
    {"type": "caseid", "name": "caseid"},
    {"type": "text audit", "name": "text_audit", "appearance": "p=100"},
    {"type": "comments", "name": "field_comments"},

    # Identification and preload
    {"type": "begin group", "name": "id_grp", "label": "Identification"},
    {"type": "select_one hhid", "name": "hhid", "label": "Household ID", "required": "yes"},
    {"type": "calculate", "name": "pl_head", "calculation": "pulldata('hh_preload', 'head_name', 'hhid', ${hhid})"},
    {"type": "calculate", "name": "pl_size", "calculation": "pulldata('hh_preload', 'hh_size', 'hhid', ${hhid})"},
    {"type": "note", "name": "pl_note", "label": "Preloaded head of household: ${pl_head} (${pl_size} members)"},
    {"type": "select_one yesno", "name": "consent", "label": "Does the respondent agree to the interview?",
     "required": "yes"},
    {"type": "end group", "name": "id_grp"},

    # Main section, only with consent
    {"type": "begin group", "name": "main", "label": "Household", "relevance": "${consent} = 1"},
    {"type": "text", "name": "resp_name", "label": "Respondent's first name (fictitious)", "required": "yes"},
    {"type": "integer", "name": "resp_age", "label": "Age", "constraint": ". >= 15 and . <= 99",
     "constraint message": "15 to 99", "required": "yes"},
    {"type": "text", "name": "phone", "label": "Phone number (10 digits starting 09)",
     "appearance": "numbers_phone", "constraint": "regex(., '^09[0-9]{8}$')"},
    {"type": "decimal", "name": "income", "label": "Monthly income", "constraint": ". >= 0"},
    {"type": "date", "name": "dob", "label": "Date of birth", "constraint": ". <= today()"},
    {"type": "time", "name": "wake_time", "label": "Usual wake-up time"},
    {"type": "datetime", "name": "next_visit", "label": "Next visit"},
    {"type": "select_one region", "name": "region", "label": "Region", "required": "yes"},
    {"type": "select_one district", "name": "district", "label": "District", "choice_filter": "filter=${region}"},
    {"type": "select_multiple assets", "name": "assets", "label": "Assets owned",
     "constraint": "not(selected(., '-99') and count-selected(.) > 1)",
     "constraint message": "'Don't know' cannot be combined"},
    {"type": "text", "name": "assets_other", "label": "Other asset", "relevance": "selected(${assets}, '96')"},
    {"type": "geopoint", "name": "gps", "label": "Household location"},
    {"type": "image", "name": "house_photo", "label": "Photo of the house (any image)"},
    {"type": "calculate", "name": "is_adult", "calculation": "if(${resp_age} >= 18, 1, 0)"},
    {"type": "calculate_here", "name": "time_main", "calculation": "once(duration())"},
    {"type": "integer", "name": "hh_size", "label": "Number of household members (1 to 3 for this test)",
     "constraint": ". >= 1 and . <= 3", "required": "yes"},

    # Counted repeat with a nested, enumerator-controlled repeat
    {"type": "begin repeat", "name": "member", "label": "Member", "repeat_count": "${hh_size}"},
    {"type": "calculate", "name": "member_index", "calculation": "index()"},
    {"type": "text", "name": "member_name", "label": "Member's first name (fictitious)", "required": "yes"},
    {"type": "select_one sex", "name": "member_sex", "label": "Sex"},
    {"type": "integer", "name": "member_age", "label": "Age", "constraint": ". >= 0 and . <= 110"},
    {"type": "select_multiple activities", "name": "member_activities", "label": "Activities"},
    {"type": "begin repeat", "name": "job", "label": "Job (add 0 to 2)"},
    {"type": "text", "name": "job_title", "label": "Job title"},
    {"type": "decimal", "name": "job_hours", "label": "Hours per week", "constraint": ". >= 0 and . <= 100"},
    {"type": "end repeat", "name": "job"},
    {"type": "end repeat", "name": "member"},
    {"type": "calculate", "name": "member_names", "calculation": "join(' ', ${member_name})"},
    {"type": "calculate", "name": "n_members", "calculation": "count(${member})"},

    # Enumerator-controlled top-level repeat
    {"type": "begin repeat", "name": "plot", "label": "Plot (add 1 or 2)"},
    {"type": "select_one crop", "name": "plot_crop", "label": "Main crop"},
    {"type": "decimal", "name": "plot_area", "label": "Area (ha)", "constraint": ". > 0 and . <= 50"},
    {"type": "end repeat", "name": "plot"},

    {"type": "text", "name": "disabled_field", "label": "Disabled field", "disabled": "yes"},
    {"type": "note", "name": "thanks", "label": "Thank you."},
    {"type": "end group", "name": "main"},
]

CHOICES = [
    ("yesno", "1", "Yes", ""), ("yesno", "0", "No", ""),
    ("region", "1", "North", ""), ("region", "2", "South", ""),
    ("district", "11", "North-A", "1"), ("district", "12", "North-B", "1"),
    ("district", "21", "South-A", "2"), ("district", "22", "South-B", "2"),
    ("assets", "1", "Radio", ""), ("assets", "2", "Bicycle", ""), ("assets", "3", "Phone", ""),
    ("assets", "96", "Other", ""), ("assets", "-99", "Don't know", ""), ("assets", "-98", "Refused", ""),
    ("sex", "1", "Male", ""), ("sex", "2", "Female", ""),
    ("activities", "1", "Farming", ""), ("activities", "2", "Trade", ""),
    ("activities", "3", "School", ""), ("activities", "-97", "None", ""),
    ("crop", "1", "Maize", ""), ("crop", "2", "Teff", ""), ("crop", "3", "Sorghum", ""),
    ("hhid", "H001", "H001", ""), ("hhid", "H002", "H002", ""), ("hhid", "H003", "H003", ""),
]

PRELOAD = [
    ("hhid", "head_name", "hh_size", "village"),
    ("H001", "Abebe (fictitious)", "2", "Village 1"),
    ("H002", "Chaltu (fictitious)", "3", "Village 2"),
    ("H003", "Dawit (fictitious)", "1", "Village 3"),
]


def main() -> None:
    wb = Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(SURVEY_HEADER)
    for row in SURVEY:
        survey.append([row.get(column) for column in SURVEY_HEADER])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "value", "label", "filter"])
    for row in CHOICES:
        choices.append([cell or None for cell in row])

    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "default_language"])
    settings.append(["ctosim coverage test", "ctosim_coverage", "2610090001", "english"])
    wb.save(HERE / "ctosim_coverage.xlsx")

    with open(HERE / "hh_preload.csv", "w", encoding="utf-8", newline="\n") as f:
        for row in PRELOAD:
            f.write(",".join(row) + "\n")


if __name__ == "__main__":
    main()
