"""Builds the third coverage XLSForm, its preload file and its search() dataset.

The first two coverage forms settled the export layout. This one covers the SurveyCTO
features and cases that no real export has shown yet (see README.md for the full list):

- an "oracle" part with no input: calculations on fixed values, and a repeat computed
  entirely by calculations, so that the export gives SurveyCTO's own result for each
  function, to compare with the simulator's;
- field types not yet exported (audio, video, file, barcode, geotrace, geoshape, sensors,
  speed violations audit is left out: its syntax is not documented);
- appearances that change what is stored (month-year, month, search(), signature);
- defaults, calculated defaults, "other (specify)", cascading and filtered selects;
- CSV edge cases: commas, quotes, accents, line breaks, leading zeros, large and small numbers;
- structure: field-list and table groups, a repeat inside a group inside a repeat, a counted
  repeat with zero instances, a repeat count given by an expression, a long field name;
- instance_name with a trailing space;
- a second version of the form (one field added), to see how an export mixes versions.

Every label says what to enter, so the expected values are known. All values are fictitious.

Run:  .venv/Scripts/python samples/coverage3/build_coverage3_form.py
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

SURVEY_HEADER = ["type", "name", "label", "hint", "default", "appearance", "constraint",
                 "constraint message", "relevance", "required", "read only", "calculation",
                 "repeat_count", "choice_filter", "minimum_seconds"]

# Calculations on fixed values: SurveyCTO's export gives the reference result of each one.
ORACLE = [
    # numbers and their formatting
    ("f_div", "10 div 4"),
    ("f_third", "1 div 3"),
    ("f_product", "2 * 3"),
    ("f_mod", "17 mod 5"),
    ("f_negative", "-7 + 2"),
    ("f_small_sum", "0.1 + 0.2"),
    ("f_big", "123456789 * 1000"),
    ("f_div_zero", "1 div 0"),
    ("f_not_a_number", "number('abc')"),
    ("f_round", "round(2.567, 2)"),
    ("f_int", "int(-2.5)"),
    ("f_pow", "pow(2, 10)"),
    ("f_sqrt", "sqrt(2)"),
    ("f_exp", "exp(1)"),
    ("f_log10", "log10(1000)"),
    ("f_pi", "pi()"),
    ("f_abs", "abs(-3.5)"),
    ("f_min_max", "concat(min(4, 2, 9), '/', max(4, 2, 9))"),
    ("f_format_number", "format-number(1234567.891)"),
    ("f_string_of_whole", "string(5.0)"),
    # logic
    ("f_true", "1 = 1"),
    ("f_false", "1 > 2"),
    ("f_if", "if(3 > 2, 'yes', 'no')"),
    ("f_coalesce", "coalesce('', 'fallback')"),
    ("f_empty", "empty('')"),
    ("f_relevant_skipped", "relevant(${never_shown})"),
    # text
    ("f_concat", "concat('a', 1, 'b')"),
    ("f_substr", "substr('abcdef', 1, 3)"),
    ("f_string_length", "string-length('héllo')"),
    ("f_lower", "lower('ÉCOLE Mixte')"),
    ("f_upper", "upper('école mixte')"),
    ("f_linebreak", "concat('line one', linebreak(), 'line two')"),
    ("f_comma_quote", "concat('a, b', ' \"quoted\"')"),
    ("f_regex", "regex('abc123', '^[a-z]+[0-9]+$')"),
    ("f_hash", "hash('abc')"),
    # dates and times
    ("f_today", "today()"),
    ("f_now", "now()"),
    ("f_format_date", "format-date-time(date('2026-01-05'), '%Y/%m/%d %a %b %e %n')"),
    ("f_format_datetime", "format-date-time(date-time('2026-01-05T14:07:09'), '%Y-%m-%d %H:%M:%S %h')"),
    ("f_date_minus_one", "date('2026-03-01') - 1"),
    ("f_date_from_number", "date(20000)"),
    ("f_date_before_today", "date('2026-01-05') < today()"),
    ("f_decimal_date_time", "decimal-date-time('2026-01-05T12:00:00')"),
    ("f_decimal_time", "decimal-time(${tm_noon})"),
    ("f_version", "version()"),
    # select and list functions
    ("f_selected", "selected('a b c', 'b')"),
    ("f_count_selected", "count-selected('a b c')"),
    ("f_selected_at", "selected-at('a b c', 1)"),
    ("f_choice_label", "choice-label(${sel_default}, ${sel_default})"),
    ("f_count_items", "count-items(' ', 'x y z')"),
    ("f_item_at", "item-at(',', 'x,y,z', 1)"),
    ("f_item_index", "item-index(',', 'x,y,z', 'z')"),
    ("f_item_present", "item-present(',', 'x,y,z', 'w')"),
    ("f_de_duplicate", "de-duplicate(' ', 'b a b c a')"),
    ("f_rank_value", "rank-value(20, '10 30 20')"),
    # randomness (format only)
    ("f_uuid", "once(uuid())"),
    ("f_random", "once(random())"),
]

SURVEY = [
    # ---- metadata, audits and sensors (Collect-only ones stay empty on the web) ----
    {"type": "start", "name": "starttime"},
    {"type": "end", "name": "endtime"},
    {"type": "deviceid", "name": "deviceid"},
    {"type": "subscriberid", "name": "subscriberid"},
    {"type": "simserial", "name": "simid"},
    {"type": "phonenumber", "name": "devicephonenum"},
    {"type": "username", "name": "username"},
    {"type": "calculate", "name": "duration", "calculation": "duration()"},
    {"type": "caseid", "name": "caseid"},
    {"type": "text audit", "name": "text_audit", "appearance": "p=100"},
    {"type": "audio audit", "name": "audio_audit", "appearance": "p=100; s=0; d=10"},
    {"type": "speed violations count", "name": "sv_count", "appearance": "100"},
    {"type": "speed violations list", "name": "sv_list", "appearance": "100"},
    {"type": "comments", "name": "field_comments"},
    {"type": "sensor_statistic mean_light_level", "name": "light_mean"},
    {"type": "sensor_statistic pct_quiet", "name": "pct_quiet"},

    {"type": "text", "name": "resp_code", "label": "Submission code: type A, B or C as told in the instructions",
     "required": "yes", "constraint": "regex(., '^[ABC]$')", "constraint message": "A, B or C",
     "minimum_seconds": "10"},

    # ---- oracle part: no input ----
    {"type": "begin group", "name": "oracle", "label": "Calculations (nothing to enter)"},
    {"type": "select_one abc", "name": "sel_default", "label": "Pre-selected answer: leave Beta selected",
     "default": "b"},
    {"type": "text", "name": "never_shown", "label": "Never shown", "relevance": "1 = 2"},
    *[{"type": "calculate", "name": name, "calculation": expr} for name, expr in ORACLE],
    {"type": "integer", "name": "n_auto", "label": "Number of computed rows (read only, 3)",
     "default": "3", "read only": "yes", "required": "yes"},
    {"type": "begin repeat", "name": "auto", "label": "Computed row", "repeat_count": "${n_auto}"},
    {"type": "calculate", "name": "a_index", "calculation": "index()"},
    {"type": "calculate", "name": "a_value", "calculation": "index() * 10"},
    {"type": "calculate", "name": "a_kind", "calculation": "if(index() mod 2 = 1, 'odd', 'even')"},
    {"type": "end repeat", "name": "auto"},
    {"type": "calculate", "name": "agg_count", "calculation": "count(${auto})"},
    {"type": "calculate", "name": "agg_sum", "calculation": "sum(${a_value})"},
    {"type": "calculate", "name": "agg_min", "calculation": "min(${a_value})"},
    {"type": "calculate", "name": "agg_max", "calculation": "max(${a_value})"},
    {"type": "calculate", "name": "agg_join", "calculation": "join(' ', ${a_value})"},
    {"type": "calculate", "name": "agg_count_if", "calculation": "count-if(${auto}, ${a_kind} = 'odd')"},
    {"type": "calculate", "name": "agg_sum_if", "calculation": "sum-if(${a_value}, ${a_kind} = 'odd')"},
    {"type": "calculate", "name": "agg_join_if", "calculation": "join-if(' ', ${a_value}, ${a_kind} = 'odd')"},
    {"type": "calculate", "name": "agg_min_if", "calculation": "min-if(${a_value}, ${a_kind} = 'odd')"},
    {"type": "calculate", "name": "agg_max_if", "calculation": "max-if(${a_value}, ${a_kind} = 'odd')"},
    {"type": "calculate", "name": "agg_indexed", "calculation": "indexed-repeat(${a_value}, ${auto}, 2)"},
    {"type": "calculate", "name": "agg_rank", "calculation": "rank-index(1, ${a_value})"},
    {"type": "calculate", "name": "agg_rank_if", "calculation": "rank-index-if(1, ${a_value}, ${a_kind} = 'even')"},
    {"type": "begin repeat", "name": "expr_rep", "label": "Row counted by an expression",
     "repeat_count": "${n_auto} - 1"},
    {"type": "calculate", "name": "e_index", "calculation": "index()"},
    {"type": "end repeat", "name": "expr_rep"},
    {"type": "end group", "name": "oracle"},

    # ---- defaults ----
    {"type": "begin group", "name": "defaults", "label": "Defaults"},
    {"type": "integer", "name": "def_keep", "label": "Pre-filled 7: do NOT change it", "default": "7"},
    {"type": "text", "name": "def_change", "label": "Pre-filled 'abc': replace it with xyz", "default": "abc"},
    {"type": "integer", "name": "calc_default", "label": "Calculated (7 x 2): try to change it to 99",
     "calculation": "${def_keep} * 2"},
    {"type": "end group", "name": "defaults"},

    # ---- text and numbers ----
    {"type": "begin group", "name": "values", "label": "Text and numbers"},
    {"type": "text", "name": "t_comma_quote", "label": "Type exactly: Hello, \"world\""},
    {"type": "text", "name": "t_accents", "label": "Type exactly: Élodie à Ouagadougou, niño"},
    {"type": "text", "name": "t_leading_zeros", "label": "Type exactly: 007"},
    {"type": "text", "name": "t_numbers", "label": "Type exactly: 012345678901", "appearance": "numbers"},
    {"type": "text", "name": "t_numbers_decimal", "label": "Type exactly: 3.50", "appearance": "numbers_decimal"},
    {"type": "decimal", "name": "d_whole", "label": "Type exactly: 5"},
    {"type": "decimal", "name": "d_trailing_zero", "label": "Type exactly: 0.10"},
    {"type": "decimal", "name": "d_negative", "label": "Type exactly: -3.75"},
    {"type": "integer", "name": "i_negative", "label": "Type exactly: -12"},
    {"type": "integer", "name": "i_large", "label": "Type exactly: 123456789"},
    {"type": "integer", "name": "i_minimum_time", "label": "Type 1 (this question has a minimum time)",
     "minimum_seconds": "10"},
    {"type": "end group", "name": "values"},

    # ---- dates and times ----
    {"type": "begin group", "name": "dates", "label": "Dates and times"},
    {"type": "date", "name": "dt_old", "label": "Enter 15 June 1950"},
    {"type": "date", "name": "dt_month_year", "label": "Enter March 2025", "appearance": "month-year"},
    {"type": "date", "name": "dt_month", "label": "Enter any month (note which)", "appearance": "month"},
    {"type": "date", "name": "dt_no_calendar", "label": "Enter 31 December 2025", "appearance": "no-calendar"},
    {"type": "time", "name": "tm_after_midnight", "label": "Enter 00:05 (5 minutes past midnight)"},
    {"type": "time", "name": "tm_noon", "label": "Enter 12:30 (half past noon)"},
    {"type": "datetime", "name": "dtm_any", "label": "Enter 1 January 2026, 23:59"},
    {"type": "end group", "name": "dates"},

    # ---- selects ----
    {"type": "begin group", "name": "selects", "label": "Choices"},
    {"type": "select_one many", "name": "s_minimal", "label": "Drop-down: choose Eleven", "appearance": "minimal"},
    {"type": "select_one agree", "name": "s_likert", "label": "Scale: choose Agree", "appearance": "likert"},
    {"type": "select_one many", "name": "s_randomized", "label": "Shuffled list: choose Ten",
     "appearance": "randomized"},
    {"type": "select_multiple many", "name": "m_many", "label": "Tick One, Ten and Twelve"},
    {"type": "select_multiple many", "name": "m_blank", "label": "Leave this question BLANK (tick nothing)"},
    {"type": "select_one comma_labels", "name": "s_comma_label", "label": "Choose 'Yes, sure'"},
    {"type": "select_one other", "name": "s_other", "label": "Choose Other"},
    {"type": "text", "name": "s_other_specify", "label": "Specify: type Bananas", "relevance": "${s_other} = 96"},
    {"type": "select_one region", "name": "region", "label": "Region: choose South"},
    {"type": "select_one district", "name": "district", "label": "District: choose South-B",
     "choice_filter": "filter = ${region}"},
    {"type": "select_multiple district", "name": "districts_multi",
     "label": "Districts of the chosen region: tick both", "choice_filter": "filter = ${region}"},
    {"type": "select_one village", "name": "village",
     "label": "Village (loaded from villages.csv, South only): choose Hawassa",
     "appearance": "search('villages', 'matches', 'region', ${region})"},
    {"type": "calculate", "name": "village_label", "calculation": "choice-label(${village}, ${village})"},
    {"type": "end group", "name": "selects"},

    # ---- groups with appearances ----
    {"type": "begin group", "name": "screen", "label": "One screen (field-list)", "appearance": "field-list"},
    {"type": "integer", "name": "fl_a", "label": "Type 1"},
    {"type": "integer", "name": "fl_b", "label": "Type 2"},
    {"type": "end group", "name": "screen"},
    {"type": "begin group", "name": "table", "label": "Table (label / list-nolabel)", "appearance": "field-list"},
    {"type": "select_one freq", "name": "tbl_header", "label": "How often?", "appearance": "label"},
    {"type": "select_one freq", "name": "tbl_water", "label": "Water: Often", "appearance": "list-nolabel"},
    {"type": "select_one freq", "name": "tbl_power", "label": "Power: Never", "appearance": "list-nolabel"},
    {"type": "end group", "name": "table"},
    {"type": "begin group", "name": "skipped", "label": "Never shown group", "relevance": "1 = 2"},
    {"type": "text", "name": "skipped_text", "label": "Never shown"},
    {"type": "end group", "name": "skipped"},
    {"type": "begin group", "name": "outer", "label": "Nested groups"},
    {"type": "begin group", "name": "middle", "label": "Middle"},
    {"type": "begin group", "name": "inner", "label": "Inner"},
    {"type": "text", "name": "deep_text", "label": "Type: deep"},
    {"type": "end group", "name": "inner"},
    {"type": "end group", "name": "middle"},
    {"type": "end group", "name": "outer"},

    # ---- repeats ----
    {"type": "integer", "name": "n_people", "label": "Number of people: type 2 (C: type 0)", "required": "yes",
     "constraint": ". >= 0 and . <= 3"},
    {"type": "begin repeat", "name": "person", "label": "Person", "repeat_count": "${n_people}"},
    {"type": "text", "name": "p_name", "label": "Name (any)"},
    {"type": "integer", "name": "p_age", "label": "Age: person 1 is 30, person 2 is 5"},
    {"type": "text", "name": "p_school", "label": "School name (any)", "relevance": "${p_age} < 18"},
    {"type": "text", "name": "this_field_name_is_longer_than_thirty_two_characters",
     "label": "Long field name: type x"},
    {"type": "begin group", "name": "health", "label": "Health"},
    {"type": "begin repeat", "name": "visit", "label": "Health visit (person 1: add 2, person 2: add 0)"},
    {"type": "date", "name": "visit_date", "label": "Visit date (any)"},
    {"type": "end repeat", "name": "visit"},
    {"type": "end group", "name": "health"},
    {"type": "end repeat", "name": "person"},

    # ---- media and geography (several are Collect only) ----
    {"type": "begin group", "name": "media", "label": "Media and locations"},
    {"type": "audio", "name": "m_audio", "label": "Record 2 seconds of audio"},
    {"type": "video", "name": "m_video", "label": "Record 2 seconds of video (or skip on the web)"},
    {"type": "file", "name": "m_file", "label": "Attach any small file (or skip)"},
    {"type": "image", "name": "m_signature", "label": "Sign here", "appearance": "signature"},
    {"type": "barcode", "name": "m_barcode", "label": "Scan any barcode (or skip on the web)"},
    {"type": "geotrace", "name": "g_trace", "label": "Record a short line of 2 to 3 points"},
    {"type": "geoshape", "name": "g_shape", "label": "Record a small shape of 3 to 4 points"},
    {"type": "end group", "name": "media"},

    {"type": "note", "name": "thanks", "label": "Thank you. Add a comment on one question before finishing "
     "(not in submission C)."},
]

CHOICES = [
    ("abc", "a", "Alpha", ""), ("abc", "b", "Beta", ""), ("abc", "c", "Gamma", ""),
    *[("many", str(i), name, "") for i, name in enumerate(
        ["One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve"],
        start=1)],
    ("agree", "1", "Strongly disagree", ""), ("agree", "2", "Disagree", ""), ("agree", "3", "Neutral", ""),
    ("agree", "4", "Agree", ""), ("agree", "5", "Strongly agree", ""),
    ("comma_labels", "1", "Yes, sure", ""), ("comma_labels", "0", "No, never", ""),
    ("other", "1", "Maize", ""), ("other", "2", "Teff", ""), ("other", "96", "Other", ""),
    ("region", "1", "North", ""), ("region", "2", "South", ""),
    ("district", "11", "North-A", "1"), ("district", "12", "North-B", "1"),
    ("district", "21", "South-A", "2"), ("district", "22", "South-B", "2"),
    # search() placeholder: the value and label columns of villages.csv
    ("village", "village_id", "village_name", ""),
    ("freq", "1", "Never", ""), ("freq", "2", "Sometimes", ""), ("freq", "3", "Often", ""),
]

VILLAGES = [
    ("village_id", "village_name", "region"),
    ("V01", "Adama", "1"),
    ("V02", "Bishoftu", "1"),
    ("V03", "Hawassa", "2"),
    ("V04", "Arba Minch", "2"),
]


# Version 2 adds one question after resp_code and lets resp_code take D (its own submission).
V2_FIELD = {"type": "text", "name": "added_in_v2", "label": "New in version 2: type v2"}
V2_RESP_CODE = {"label": "Submission code: type D (version 2)",
                "constraint": "regex(., '^[ABCD]$')", "constraint message": "A, B, C or D"}


def write_form(path: Path, version: str, rows: list[dict]) -> None:
    wb = Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(SURVEY_HEADER)
    for row in rows:
        survey.append([row.get(column) for column in SURVEY_HEADER])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "value", "label", "filter"])
    for row in CHOICES:
        choices.append([cell or None for cell in row])

    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "default_language", "instance_name"])
    settings.append(["ctosim coverage test 3", "ctosim_coverage3", version, "english",
                     "concat('C3-', ${resp_code}, ' ')"])
    wb.save(path)


def main() -> None:
    write_form(HERE / "ctosim_coverage3.xlsx", "2610120001", SURVEY)
    position = next(i for i, row in enumerate(SURVEY) if row.get("name") == "resp_code") + 1
    v2 = [dict(row, **V2_RESP_CODE) if row.get("name") == "resp_code" else row for row in SURVEY]
    write_form(HERE / "ctosim_coverage3_v2.xlsx", "2610120003", v2[:position] + [V2_FIELD] + v2[position:])

    with open(HERE / "villages.csv", "w", encoding="utf-8", newline="\n") as f:
        for row in VILLAGES:
            f.write(",".join(row) + "\n")


if __name__ == "__main__":
    main()
