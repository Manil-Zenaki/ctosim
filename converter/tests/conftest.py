"""Small XLSForm fixtures derived exclusively from spec 02."""

import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook


@pytest.fixture
def workbook_factory(tmp_path):
    def make(survey=None, *, choices=None, settings=None, filename="form.xlsx",
             extra_sheets=None, omit_survey=False):
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = {}
        if not omit_survey:
            sheets["survey"] = survey if survey is not None else [
                ["type", "name", "label"], ["text", "answer", "Answer"]]
        if choices is not None:
            sheets["choices"] = choices
        if settings is not None:
            sheets["settings"] = settings
        sheets.update(extra_sheets or {})
        for name, rows in sheets.items():
            sheet = workbook.create_sheet(name)
            for row in rows:
                sheet.append(row)
        if not sheets:
            workbook.create_sheet("other")
        path = tmp_path / filename
        workbook.save(path)
        workbook.close()
        return path
    return make


@pytest.fixture
def set_formula_cache():
    """Set OOXML cached values without requiring Excel or production helpers."""
    def patch(path, sheet_number, address, value):
        member = f"xl/worksheets/sheet{sheet_number}.xml"
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with ZipFile(path) as archive:
            contents = {item.filename: archive.read(item.filename)
                        for item in archive.infolist()}
        root = ET.fromstring(contents[member])
        cell = root.find(f".//{namespace}c[@r='{address}']")
        assert cell is not None and cell.find(f"{namespace}f") is not None
        cached = cell.find(f"{namespace}v")
        if cached is None:
            cached = ET.SubElement(cell, f"{namespace}v")
        if isinstance(value, str):
            cell.set("t", "str")
        else:
            cell.attrib.pop("t", None)
        cached.text = str(value)
        contents[member] = ET.tostring(root, encoding="utf-8")
        with ZipFile(path, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)
    return patch
