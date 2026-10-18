"""Convert SurveyCTO XLSForms to XForms. See specs/02-xlsform-converter.md."""

from __future__ import annotations

import contextlib
import io
import json
import logging
import os
import re
import tempfile
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from ._normalize import CTOSIM_NAMESPACE, PASSED_SETTINGS, NormalizationError, Sheet, is_empty, normalize, read_sheet, text


class ConversionError(Exception):
    """Raised when a spreadsheet form cannot be converted."""


@dataclass(frozen=True)
class FieldInfo:
    name: str
    type: str
    xpath: str
    row: int


@dataclass(frozen=True)
class ConversionResult:
    xform: str
    form_id: str
    title: str
    version: str | None
    fields: tuple[FieldInfo, ...]
    warnings: tuple[str, ...]

    def write(self, xform_path: str | os.PathLike, catalog_path: str | os.PathLike | None = None) -> None:
        xform_path = Path(xform_path)
        catalog_path = xform_path.with_suffix(".fields.json") if catalog_path is None else Path(catalog_path)
        with open(xform_path, "w", encoding="utf-8", newline="") as f:
            f.write(self.xform)
        catalog = {
            "form_id": self.form_id,
            "title": self.title,
            "version": self.version,
            "fields": [{"name": f.name, "type": f.type, "xpath": f.xpath, "row": f.row} for f in self.fields],
        }
        with open(catalog_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(json.dumps(catalog, indent=2, ensure_ascii=False) + "\n")


def convert_xlsform(path: str | os.PathLike) -> ConversionResult:
    path = Path(os.fspath(path))
    source = path.name

    with _quiet_libraries():
        survey, choices, settings = _read_workbook(path)
    if survey is None:
        raise ConversionError(f"{source}: the workbook has no 'survey' sheet")
    try:
        norm = normalize(survey, choices, settings)
    except NormalizationError as e:
        raise ConversionError(f"{source}: {e}") from None

    form_id = text(norm.settings.get("form_id")).strip() or _form_id_from_file(path)
    title = text(norm.settings.get("form_title")).strip() or form_id
    version = _version(norm.settings.get("version"))

    with _quiet_libraries():
        xform = _run_pyxform(norm, form_id, title, version, source)
    return ConversionResult(
        xform=xform,
        form_id=form_id,
        title=title,
        version=version,
        fields=tuple(FieldInfo(*f) for f in norm.fields),
        warnings=tuple(norm.warnings),
    )


@contextlib.contextmanager
def _quiet_libraries():
    """Keeps the messages of the libraries the converter uses off stderr: only the converter's own
    warnings are written (spec 02, Amendment 6)."""
    previous = logging.root.manager.disable
    logging.disable(logging.CRITICAL)
    try:
        with warnings.catch_warnings(), contextlib.redirect_stderr(io.StringIO()):
            warnings.simplefilter("ignore")
            yield
    finally:
        logging.disable(previous)


def _read_workbook(path: Path) -> tuple[Sheet | None, Sheet | None, Sheet | None]:
    import openpyxl

    try:
        # data_only: a formula cell yields the value Excel cached with the file, or None.
        workbook = openpyxl.load_workbook(path, data_only=True)
    except OSError as e:
        raise ConversionError(f"{path.name}: the file cannot be read ({e})") from e
    except Exception as e:
        raise ConversionError(f"{path.name}: the file is not an .xlsx workbook ({e})") from e
    try:
        sheets = {name: read_sheet(workbook[name]) for name in ("survey", "choices", "settings")
                  if name in workbook.sheetnames}
    finally:
        workbook.close()
    return sheets.get("survey"), sheets.get("choices"), sheets.get("settings")


def _form_id_from_file(path: Path) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]", "_", path.stem)
    return stem if re.match(r"[A-Za-z]", stem) else "f_" + stem


def _version(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _run_pyxform(norm, form_id: str, title: str, version: str | None, source: str) -> str:
    from pyxform.errors import PyXFormError
    from pyxform.xls2xform import convert

    with tempfile.TemporaryDirectory(prefix="ctosim-") as tmp:
        workbook_path = Path(tmp) / "form.xlsx"
        rows = _write_odk_workbook(norm, form_id, title, version, workbook_path)
        try:
            with warnings.catch_warnings(), _repeats_compared_by_identity(), _typographic_quotes_kept():
                warnings.simplefilter("ignore")
                result = convert(xlsform=str(workbook_path), validate=False, pretty_print=True,
                                 warnings=[])
        except PyXFormError as e:
            raise ConversionError(f"{source}: {_with_row(_user_rows(str(e), rows), norm)}") from e
        except Exception as e:
            raise ConversionError(f"{source}: the form could not be converted ({e})") from e
    return _exact_version_attribute(result.xform, version)


@contextlib.contextmanager
def _repeats_compared_by_identity():
    """
    pyxform 4.5 compares repeat sections with ==, which serialises both whole sections each time:
    a form with many questions in repeats took minutes. The comparison there asks whether two
    sections are the same one, so it is made by identity while pyxform runs; the XForm is unchanged.
    """
    from pyxform.section import RepeatingSection

    had_eq, had_hash = "__eq__" in RepeatingSection.__dict__, "__hash__" in RepeatingSection.__dict__
    old_eq, old_hash = RepeatingSection.__dict__.get("__eq__"), RepeatingSection.__dict__.get("__hash__")
    RepeatingSection.__eq__ = lambda self, other: self is other
    RepeatingSection.__hash__ = object.__hash__
    try:
        yield
    finally:
        for name, had, old in (("__eq__", had_eq, old_eq), ("__hash__", had_hash, old_hash)):
            if had:
                setattr(RepeatingSection, name, old)
            else:
                delattr(RepeatingSection, name)


@contextlib.contextmanager
def _typographic_quotes_kept():
    """
    pyxform replaces typographic quotes with straight ones in every cell, which ends a string
    literal early in 'l’argent'; SurveyCTO keeps them as written (spec 02, Amendment 7).
    """
    from pyxform.parsing import sheet_headers

    old = sheet_headers.RE_SMART_QUOTES
    sheet_headers.RE_SMART_QUOTES = re.compile(r"(?!)")  # matches nothing
    try:
        yield
    finally:
        sheet_headers.RE_SMART_QUOTES = old


def _exact_version_attribute(xform: str, version: str | None) -> str:
    """pyxform trims the version; the instance root must carry it exactly as stored."""
    if version is None or version == version.strip():
        return xform
    root = re.search(r"<instance>\s*<data\b[^>]*>", xform)
    if root is None:
        return xform
    tag = root.group(0)
    fixed = re.sub(r'\bversion="[^"]*"', lambda _: f'version="{escape(version, {chr(34): "&quot;"})}"', tag, count=1)
    return xform[: root.start()] + fixed + xform[root.end():]


def _user_rows(message: str, rows: dict[str, list[int]]) -> str:
    """pyxform numbers the rows of the workbook it reads, which has no empty rows: give the user's."""
    sheet_rows = rows["choices"] if re.search(r"\bchoices\b", message, re.IGNORECASE) and "choices" in rows \
        else rows["survey"]

    def user(m: re.Match) -> str:
        n = int(m.group(2))
        return m.group(1) + str(sheet_rows[n] if 0 <= n < len(sheet_rows) and sheet_rows[n] else n)

    return re.sub(r"(\brow\s*:?\s*)(\d+)", user, message, flags=re.IGNORECASE)


def _with_row(message: str, norm) -> str:
    """Adds the survey row to a pyxform message that names a type or field but no row."""
    if re.search(r"\brow\b", message, re.IGNORECASE):
        return message
    t_col, n_col = norm.survey.column("type"), norm.survey.column("name")
    for quoted in re.findall(r"'([^']+)'", message):
        for k, row in enumerate(norm.survey.rows):
            if quoted in (text(row[t_col]).strip(), text(row[n_col]).strip()):
                return f"row {k + 2}: {message}"
    return message


def _write_odk_workbook(norm, form_id: str, title: str, version: str | None, out: Path) -> dict[str, list[int]]:
    """Writes the workbook pyxform reads; returns, per sheet, the user's row number of each written row."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rows = {"survey": _write_sheet(wb.create_sheet("survey"), norm.survey)}
    if norm.choices is not None:
        rows["choices"] = _write_sheet(wb.create_sheet("choices"), norm.choices)

    settings = wb.create_sheet("settings")
    headers = ["form_id", "form_title", "allow_choice_duplicates"]
    values: list[Any] = [form_id, title, "yes"]
    if version is not None:
        headers.append("version")
        values.append(version)
    for name in PASSED_SETTINGS:
        value = norm.settings.get(name)
        if value is not None and text(value).strip() != "":
            headers.append(name)
            values.append(value)
    if norm.uses_ctosim_namespace:
        own = text(norm.settings.get("namespaces")).strip()
        headers.append("namespaces")
        values.append(f"{own} {CTOSIM_NAMESPACE}".strip())
    _append(settings, headers)
    _append(settings, values)
    wb.save(out)
    return rows


def _write_sheet(ws, sheet: Sheet) -> list[int]:
    """
    Writes the header and the non-empty rows: pyxform stops reading after a run of empty rows, so
    empty and left-out rows are not written. Returns the user's row number of each written row
    (index = row number in the written sheet; 0 for unused positions).
    """
    # pyxform reads 'label:French' as a translation only when no header uses '::' (as the
    # bind:: columns this converter adds do), so translated and media columns are written with '::'.
    _append(ws, [h if "::" in h or ":" not in h else h.replace(":", "::", 1) for h in sheet.headers])
    user_rows = [0, 1]
    for k, row in enumerate(sheet.rows):
        if is_empty(row):
            continue
        _append(ws, row)
        user_rows.append(k + 2)
        for added in sheet.added_after.get(k, ()):
            _append(ws, added)
            user_rows.append(k + 2)
    return user_rows


def _append(ws, values: list[Any]) -> None:
    ws.append([None if v == "" else v for v in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"  # literal text, not a formula


__all__ = ["ConversionError", "ConversionResult", "FieldInfo", "convert_xlsform"]
